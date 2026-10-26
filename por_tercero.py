"""Agrupa un FORMATO_REEMBOLSO / CAJA_MENOR por tercero (USD). Uso: python3 -I por_tercero.py ENTRADA.xlsx SALIDA.xlsx "TITULO" """
import sys, re, datetime as dt
from collections import OrderedDict
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def num(v):
    if v is None or str(v).strip() == '': return None
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip()
    if s.count(',') > 1: s = s.replace(',', '', s.count(',') - 1)
    return float(s.replace(',', '.'))

def fecha(v):
    if isinstance(v, dt.datetime): return v
    s = str(v or '').strip().replace('//', '/')
    m = re.fullmatch(r'(\d{1,2})/?(\d{2})/(\d{4})', s)   # '2102/2025' -> 21/02/2025
    return dt.datetime(int(m[3]), int(m[2]), int(m[1])) if m else (v or None)

src, out, titulo = sys.argv[1:4]
wb = load_workbook(src); o = wb['CAJAS REEMBOLSO ']
g = OrderedDict()
for r in range(9, o.max_row + 1):
    j, k = num(o.cell(r, 10).value), num(o.cell(r, 11).value)
    m = o.cell(r, 13).value
    m = None if isinstance(m, str) and m.startswith('=') else num(m)
    if o.cell(r, 1).value is None or (j is None and k is None and m is None): continue
    j, k = j or 0, k or 0
    m = (j + k) if m is None else m
    t = (str(o.cell(r, 6).value or '').strip().upper()) or '(SIN TERCERO)'
    g.setdefault(t, []).append((o.cell(r, 4).value, fecha(o.cell(r, 5).value), o.cell(r, 7).value, j, k, m, o.cell(r, 1).value))

ws = wb.create_sheet('Por tercero', 0)
hf = PatternFill('solid', fgColor='1F3864'); pf = PatternFill('solid', fgColor='D9E1F2')
ws['A1'] = titulo; ws['A1'].font = Font(bold=True, size=13)
for i, h in enumerate(['Tercero', 'N° factura', 'Fecha', 'Concepto', 'Subtotal sin IVA', 'IVA', 'Total', 'Tipo', 'Ítem'], 1):
    c = ws.cell(4, i, h); c.fill = hf; c.font = Font(bold=True, color='FFFFFF'); c.alignment = Alignment(horizontal='center')
r = 5
for t, fs in sorted(g.items()):
    p = r; a, b = r + 1, r + len(fs)
    ws.cell(p, 1, t); ws.cell(p, 8, 'TOTAL TERCERO')
    ws.cell(p, 2, f'=COUNTIFS($A${a}:$A${b},$A{p},$H${a}:$H${b},"Factura")').number_format = '0" factura(s)"'
    for col in 'EFG':
        ws[f'{col}{p}'] = f'=SUM({col}{a}:{col}{b})'
    for c in range(1, 10):
        x = ws.cell(p, c); x.fill = pf; x.font = Font(bold=True)
    r += 1
    for n, f, cpt, j, k, m, it in fs:
        for c, v in enumerate([t, n, f, cpt, j, k, m, 'Factura', it], 1): ws.cell(r, c, v)
        ws.cell(r, 3).number_format = 'dd/mm/yyyy'
        r += 1
last = r - 1
ws['D2'] = 'TOTAL GENERAL'; ws['D2'].font = Font(bold=True)
for col in 'EFG':
    ws[f'{col}2'] = f'=SUMIF($H$5:$H${last},"TOTAL TERCERO",{col}5:{col}{last})'; ws[f'{col}2'].font = Font(bold=True)
ws['B2'] = f'=COUNTIF($H$5:$H${last},"Factura")'; ws['B2'].number_format = '0" facturas"'; ws['B2'].font = Font(bold=True)
for row in ws.iter_rows(min_row=2, max_row=last, min_col=5, max_col=7):
    for c in row: c.number_format = '#,##0.00'
for col, w in zip('ABCDEFGHI', [28, 24, 12, 34, 16, 12, 14, 16, 7]): ws.column_dimensions[col].width = w
ws.freeze_panes = 'B5'; ws.auto_filter.ref = f'A4:I{last}'
wb.save(out)
