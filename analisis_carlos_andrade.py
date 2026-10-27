import openpyxl, datetime as dt
from openpyxl.styles import Font, PatternFill, Alignment
src='/root/.claude/uploads/68572e23-a5b0-58f8-97dc-3995dde3fc9e/e303380d-Libro1.xlsx'
out='salidas/Libro1_carlos_andrade_analisis.xlsx'
wb=openpyxl.load_workbook(src); S="'carlos andrade'!"
ws=wb.create_sheet('Resumen mensual',0)
hf=PatternFill('solid',fgColor='1F3864'); ok=PatternFill('solid',fgColor='C6E0B4'); bad=PatternFill('solid',fgColor='F8CBAD')
ws['A1']='Cuenta 13659501 OTROS — Carlos Andrade: movimiento y saldo por mes'; ws['A1'].font=Font(bold=True,size=13)
for i,h in enumerate(['Mes','Débitos (cobros)','Créditos (descuentos)','Neto del mes','Saldo a fin de mes','Estado','Cambio vs mes anterior'],1):
    c=ws.cell(3,i,h); c.fill=hf; c.font=Font(bold=True,color='FFFFFF'); c.alignment=Alignment(horizontal='center',wrap_text=True)
r=4; y,m=2022,11
while (y,m)<=(2026,9):
    ws.cell(r,1,dt.datetime(y,m,1)).number_format='mmm-yyyy'
    lo=f'$A{r}'; hi=f'EDATE($A{r},1)'
    for col,sc in ((2,'J'),(3,'K')):
        ws.cell(r,col,f'=SUMIFS({S}${sc}$7:${sc}$229,{S}$A$7:$A$229,">="&{lo},{S}$A$7:$A$229,"<"&{hi})')
    ws.cell(r,4,f'=B{r}-C{r}'); ws.cell(r,5,f'=D{r}' if r==4 else f'=E{r-1}+D{r}')
    ws.cell(r,6,f'=IF(E{r}=0,"Cuadrado en 0","Diferencia")'); ws.cell(r,7,f'=E{r}' if r==4 else f'=E{r}-E{r-1}')
    for c in range(2,8):
        if c!=6: ws.cell(r,c).number_format='#,##0;[Red]-#,##0'
    r+=1; m+=1
    if m==13: y+=1; m=1
last=r-1
from openpyxl.formatting.rule import FormulaRule
ws.conditional_formatting.add(f'F4:F{last}',FormulaRule(formula=['$E4=0'],fill=ok))
ws.conditional_formatting.add(f'F4:F{last}',FormulaRule(formula=['$E4<>0'],fill=bad))
ws.cell(last+1,1,'Total').font=Font(bold=True)
for c in 'BCD': ws[f'{c}{last+1}']=f'=SUM({c}4:{c}{last})'; ws[f'{c}{last+1}'].number_format='#,##0'; ws[f'{c}{last+1}'].font=Font(bold=True)
ws[f'E{last+1}']=f'=E{last}'; ws[f'E{last+1}'].number_format='#,##0;[Red]-#,##0'; ws[f'E{last+1}'].font=Font(bold=True)
ws[f'F{last+1}']=f'=IF(ROUND(E{last+1}-{S}L229,0)=0,"✔ cuadra con el auxiliar","revisar")'
for col,w in zip('ABCDEFG',[12,18,20,16,18,18,20]): ws.column_dimensions[col].width=w
ws.freeze_panes='A4'
# Puente
p=wb.create_sheet('Por qué 94.827',1)
p['A1']='¿De dónde salen los $94.827 al 30/09/2026?'; p['A1'].font=Font(bold=True,size=13)
rows=[('Saldo a 31/07/2026 (antes del ajuste)',f"='Resumen mensual'!E{4+ (2026-2022)*12+7-11}",'Descuentos de póliza de vehículo/exequial superiores a lo amortizado/facturado desde sep-2025'),
('+ NT-12625 del 31/08/26 "Corrección tercero NTE 12088"',f'=SUM({S}J216:J221)','Débito: amortizaciones (35.700, 131.733, 142.268, cada una x2) pasadas a este tercero'),
('+ FSC-11290 del 27/08/26 póliza moto (62.590) cobrada y no descontada en agosto',f'={S}J215','En agosto la quincena solo descontó 117.897 (carro + exequial), sin la moto. En septiembre se descontó la de septiembre'),
('± Otros movimientos de ago-sep 2026 (septiembre cobro 190.990 vs descuento 190.991, etc.)',None,'Diferencia de redondeo'),
('− NF-355 del 30/09/26 descuento póliza de salud (feb-mar-abr)',f'=-SUM({S}K227:K229)','Las facturas de salud (FSC-9538 y 9539, 91.984) ya habían sido anuladas con las NC de abril (NT-12434 y 12435). Este descuento queda sin cobro que lo respalde'),
('Saldo al 30/09/2026 (resultado)',f'={S}L229','')]
for i,h in enumerate(['Concepto','Valor','Qué pasó'],1):
    c=p.cell(3,i,h); c.fill=hf; c.font=Font(bold=True,color='FFFFFF')
for i,(a,b,c) in enumerate(rows,4):
    p.cell(i,1,a); p.cell(i,2,b); p.cell(i,3,c); p.cell(i,2).number_format='#,##0;[Red]-#,##0'; p.cell(i,3).alignment=Alignment(wrap_text=True,vertical='top'); p.cell(i,1).alignment=Alignment(wrap_text=True,vertical='top')
p['B7']='=B9-B4-B5-B6-B8'
p['A9'].font=p['B9'].font=Font(bold=True)
p['A11']='Reparto del resultado (sin el descuento de salud): 186.810 corresponden a vehículo/exequial y −91.983 al descuento de salud sin cobro.'
p.column_dimensions['A'].width=60; p.column_dimensions['B'].width=16; p.column_dimensions['C'].width=80
wb.save(out)
